from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


def write_excel(df: pd.DataFrame, path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="Comparativa", index=False)
        ws = writer.sheets["Comparativa"]
        header_fill = PatternFill("solid", fgColor="1F4E78")
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
        for idx, col in enumerate(df.columns, start=1):
            letter = get_column_letter(idx)
            width = max(len(str(col)), *(len(str(v)) for v in df[col])) if len(df) else len(str(col))
            ws.column_dimensions[letter].width = min(width + 2, 45)
            for cell in ws[letter][1:]:
                if col == "Margen %":
                    cell.number_format = "0.0%"
                elif df[col].dtype.kind == "f":
                    cell.number_format = "#,##0"
        ws.freeze_panes = "B2"
    return path
